"""
Session logger.

  bubbles.csv   appended + flushed on every bubble (survives a crash / power cut)
  bubbles.xlsx  rebuilt on demand and every few seconds.  Diameter in mm and sphere volume are
                Excel FORMULAS, driven by the "mm per pixel" cell on the Settings sheet, so
                you can change the calibration in Excel and everything recomputes.

Volume model: each bubble is a sphere of its measured diameter,  V = pi/6 * d^3.
(1 mm^3 = 1 microlitre; 1000 mm^3 = 1 mL.)
"""
import csv
import math
import os
from datetime import datetime, timedelta
from pathlib import Path

CSV_FIELDS = ['id', 'time_s', 'clock_time', 'diameter_px', 'major_px', 'minor_px', 'method', 'confidence']


def sphere_volume_mm3(d_mm):
    return math.pi / 6.0 * d_mm ** 3


class SessionLogger:
    def __init__(self, root_dir, mm_per_px=0.03, label=''):
        self.start = datetime.now()
        stamp = self.start.strftime('%Y%m%d_%H%M%S')
        self.dir = Path(root_dir).expanduser() / f'session_{stamp}'
        self.dir.mkdir(parents=True, exist_ok=True)
        self.csv_path = self.dir / 'bubbles.csv'
        self.xlsx_path = self.dir / 'bubbles.xlsx'
        self.mm_per_px = float(mm_per_px)
        self.label = label
        self.rows = []
        self.dirty = False
        self.duration_s = 0.0
        self._fh = open(self.csv_path, 'w', newline='')
        self._w = csv.DictWriter(self._fh, fieldnames=CSV_FIELDS)
        self._w.writeheader()
        self._fh.flush()

    # ------------------------------------------------------------------ data
    def add(self, rec):
        t = float(rec['t_last'])
        row = dict(id=rec['id'], time_s=round(t, 2),
                   clock_time=(self.start + timedelta(seconds=t)).strftime('%Y-%m-%d %H:%M:%S'),
                   diameter_px=round(float(rec['d_px']), 2), major_px=round(float(rec['major_px']), 2),
                   minor_px=round(float(rec['minor_px']), 2), method=rec['method'],
                   confidence='low' if rec['low_conf'] else 'ok')
        self.rows.append(row)
        self._w.writerow(row)
        self._fh.flush()
        self.dirty = True
        self.duration_s = max(self.duration_s, t)

    def set_duration(self, seconds):
        self.duration_s = max(self.duration_s, float(seconds))

    def set_mm_per_px(self, v):
        self.mm_per_px = float(v)
        self.dirty = True

    def stats(self, include_low=True, elapsed_s=None):
        rows = [r for r in self.rows if include_low or r['confidence'] == 'ok']
        d = [r['diameter_px'] * self.mm_per_px for r in rows]
        vol = sum(sphere_volume_mm3(x) for x in d)
        dur = self.duration_s if elapsed_s is None else elapsed_s
        top = max((sphere_volume_mm3(x) for x in d), default=0.0)
        return dict(count=len(rows), total_mm3=vol, total_ml=vol / 1000.0,
                    top_share=(top / vol) if vol > 0 else 0.0,
                    mean_mm=(sum(d) / len(d)) if d else 0.0, max_mm=max(d) if d else 0.0,
                    ml_per_min=(vol / 1000.0) / (dur / 60.0) if dur > 1 else 0.0)

    # ------------------------------------------------------------------ excel
    def save_xlsx(self, duration_s=None):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        if duration_s is not None:
            self.set_duration(duration_s)
        wb = Workbook()
        bold = Font(bold=True)
        head_fill = PatternFill('solid', fgColor='DDEBF7')

        # ---- Settings
        st = wb.active
        st.title = 'Settings'
        st['A1'], st['B1'] = 'mm per pixel (calibration)', self.mm_per_px
        st['A2'], st['B2'] = 'Session started', self.start.strftime('%Y-%m-%d %H:%M:%S')
        st['A3'], st['B3'] = 'Source', self.label
        st['A5'] = 'Change B1 and every diameter / volume in this workbook recomputes.'
        st['A6'] = 'Volume assumes each bubble is a sphere: V = pi/6 * d^3.  Volume error is ~3x the diameter error.'
        for c in ('A1', 'A2', 'A3'):
            st[c].font = bold
        st['B1'].fill = PatternFill('solid', fgColor='FFF2CC')
        st.column_dimensions['A'].width = 34
        st.column_dimensions['B'].width = 28

        # ---- Bubbles
        ws = wb.create_sheet('Bubbles')
        heads = ['ID', 'Time (s)', 'Clock time', 'Diameter (px)', 'Diameter (mm)', 'Volume (mm3)',
                 'Major axis (px)', 'Minor axis (px)', 'Method', 'Confidence']
        ws.append(heads)
        for c in ws[1]:
            c.font, c.fill = bold, head_fill
            c.alignment = Alignment(horizontal='center', wrap_text=True)
        for n, r in enumerate(self.rows, start=2):
            ws.append([r['id'], r['time_s'], r['clock_time'], r['diameter_px'],
                       f'=D{n}*Settings!$B$1', f'=PI()/6*E{n}^3',
                       r['major_px'], r['minor_px'], r['method'], r['confidence']])
            ws[f'E{n}'].number_format = '0.000'
            ws[f'F{n}'].number_format = '0.0000'
        for col, w in zip('ABCDEFGHIJ', (7, 10, 20, 14, 14, 14, 15, 15, 16, 12)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = 'A2'

        # ---- Summary (all formulas)
        sm = wb.create_sheet('Summary', 0)
        rng = lambda c: f'Bubbles!{c}2:{c}200000'      # generous fixed range, grows with the log
        items = [
            ('Bubbles logged', f'=COUNT({rng("A")})', '0'),
            ('Total volume (mm3 = microlitres)', f'=SUM({rng("F")})', '0.000'),
            ('Total volume (mL)', '=B3/1000', '0.00000'),
            ('Total volume, excluding low-confidence bubbles (mm3)',
             f'=SUMIFS({rng("F")},{rng("J")},"ok")', '0.000'),
            ('Mean diameter (mm)', f'=IFERROR(AVERAGE({rng("E")}),0)', '0.000'),
            ('Median diameter (mm)', f'=IFERROR(MEDIAN({rng("E")}),0)', '0.000'),
            ('Largest diameter (mm)', f'=IFERROR(MAX({rng("E")}),0)', '0.000'),
            ('Session length (s)', round(self.duration_s, 1), '0.0'),
            ('Gas rate (mL / min)', '=IF(B9>0,B4/(B9/60),0)', '0.0000'),
            ('Gas rate (mL / hour)', '=B10*60', '0.000'),
            ('Largest single bubble: share of total volume', f'=IFERROR(MAX({rng("F")})/B3,0)', '0%'),
        ]
        sm['A1'], sm['B1'] = 'BubbleMeter session summary', ''
        sm['A1'].font = Font(bold=True, size=13)
        for k, (lab, val, fmt) in enumerate(items, start=2):
            sm[f'A{k}'], sm[f'B{k}'] = lab, val
            sm[f'A{k}'].font = bold
            sm[f'B{k}'].number_format = fmt
        sm['A14'] = ('Note: bubbles under ~8 px are near the noise floor ("low" confidence); ring-style '
                     'bubbles can be under-measured. Volume scales with diameter cubed, so check the largest '
                     'bubbles on the Bubbles sheet (sort by Volume) - one bad detection can dominate the total.')
        sm.column_dimensions['A'].width = 52
        sm.column_dimensions['B'].width = 18

        tmp = str(self.xlsx_path) + '.tmp'
        wb.save(tmp)
        os.replace(tmp, self.xlsx_path)
        self.dirty = False
        return self.xlsx_path

    def close(self):
        try:
            self._fh.close()
        except Exception:
            pass
